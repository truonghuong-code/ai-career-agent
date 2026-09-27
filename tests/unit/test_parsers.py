from io import BytesIO

from docx import Document

from app.rag.parsers import DocumentParserRegistry


def test_text_and_markdown_parser_extract_utf8_content() -> None:
    registry = DocumentParserRegistry()

    _, text_parser = registry.get_parser("notes.txt", "text/plain")
    markdown_type, markdown_parser = registry.get_parser("notes.md", "text/markdown")

    assert text_parser.extract(b"Python and PostgreSQL").text == "Python and PostgreSQL"
    assert markdown_parser.extract(b"# Career plan").text == "# Career plan"
    assert markdown_type == "markdown"


def test_docx_parser_extracts_paragraphs() -> None:
    document = Document()
    document.add_paragraph("AI Engineer")
    document.add_paragraph("PostgreSQL and pgvector")
    buffer = BytesIO()
    document.save(buffer)

    _, parser = DocumentParserRegistry().get_parser(
        "profile.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    extracted = parser.extract(buffer.getvalue())

    assert extracted.text == "AI Engineer\nPostgreSQL and pgvector"
    assert extracted.metadata == {"paragraph_count": 2}


def test_pdf_parser_is_selected_by_extension() -> None:
    document_type, _ = DocumentParserRegistry().get_parser("document.pdf", "")

    assert document_type == "pdf"
