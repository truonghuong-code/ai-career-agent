from io import BytesIO

import httpx
import pytest
from docx import Document as DocxDocument

from app.api.v1.endpoints.documents import get_search_service
from app.db.repositories.documents import DocumentRepository, RetrievedChunk
from app.integrations.storage.local import LocalDocumentStorage
from app.main import create_app
from app.rag.chunking import DeterministicTextChunker
from app.rag.embeddings import DeterministicEmbeddingProvider
from app.rag.parsers import DocumentParserRegistry
from app.services.document_service import DocumentIngestionService, DocumentSearchService


def _create_ingestion_service(
    db_session, storage_dir, provider: DeterministicEmbeddingProvider
) -> DocumentIngestionService:
    return DocumentIngestionService(
        repository=DocumentRepository(db_session),
        parsers=DocumentParserRegistry(),
        chunker=DeterministicTextChunker(chunk_size=200, chunk_overlap=20),
        embedding_provider=provider,
        storage=LocalDocumentStorage(storage_dir),
    )


def _create_search_service(
    db_session, provider: DeterministicEmbeddingProvider
) -> DocumentSearchService:
    return DocumentSearchService(DocumentRepository(db_session), provider)


async def test_rm_001_retrieved_chunk_exposes_source_and_chunk_contract(
    db_session, storage_dir
) -> None:
    provider = DeterministicEmbeddingProvider(dimensions=256)
    ingestion_service = _create_ingestion_service(db_session, storage_dir, provider)
    ingested = await ingestion_service.ingest(
        owner_id="user-a",
        filename="contract.txt",
        mime_type="text/plain",
        content=b"FastAPI PostgreSQL retrieval contract",
    )

    matches = await _create_search_service(db_session, provider).search(
        owner_id="user-a",
        query="FastAPI PostgreSQL retrieval contract",
        limit=5,
        document_ids=[ingested.document.id],
    )

    assert len(matches) == 1
    match = matches[0]
    assert isinstance(match, RetrievedChunk)
    assert match.chunk.id is not None
    assert match.chunk.document_id == ingested.document.id
    assert match.chunk.chunk_index == 0
    assert match.chunk.text == "FastAPI PostgreSQL retrieval contract"
    assert match.chunk.chunk_metadata == {}
    assert match.chunk.document.filename == "contract.txt"
    assert match.chunk.document.document_type == "txt"
    assert match.score == pytest.approx(1.0)


async def test_rm_002_retrieval_score_is_similarity_with_higher_match_first(
    db_session, storage_dir
) -> None:
    provider = DeterministicEmbeddingProvider(dimensions=256)
    ingestion_service = _create_ingestion_service(db_session, storage_dir, provider)
    exact = await ingestion_service.ingest(
        owner_id="user-a",
        filename="exact.txt",
        mime_type="text/plain",
        content=b"FastAPI PostgreSQL vector retrieval",
    )
    less_relevant = await ingestion_service.ingest(
        owner_id="user-a",
        filename="less-relevant.txt",
        mime_type="text/plain",
        content=b"Japanese client meeting requirement clarification",
    )

    matches = await _create_search_service(db_session, provider).search(
        owner_id="user-a",
        query="FastAPI PostgreSQL vector retrieval",
        limit=5,
        document_ids=[exact.document.id, less_relevant.document.id],
    )

    assert [match.chunk.document_id for match in matches] == [
        exact.document.id,
        less_relevant.document.id,
    ]
    assert matches[0].score > matches[1].score
    assert matches[0].score == pytest.approx(1.0)


async def test_rm_003_search_api_returns_explicit_result_schema(db_session, storage_dir) -> None:
    provider = DeterministicEmbeddingProvider(dimensions=256)
    ingestion_service = _create_ingestion_service(db_session, storage_dir, provider)
    document = DocxDocument()
    document.add_paragraph("FastAPI retrieval")
    document.add_paragraph("PostgreSQL pgvector")
    content = BytesIO()
    document.save(content)
    ingested = await ingestion_service.ingest(
        owner_id="user-a",
        filename="profile.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=content.getvalue(),
    )
    search_service = _create_search_service(db_session, provider)
    expected_match = (
        await search_service.search(
            owner_id="user-a",
            query="FastAPI retrieval PostgreSQL pgvector",
            limit=5,
            document_ids=[ingested.document.id],
        )
    )[0]
    app = create_app()
    app.dependency_overrides[get_search_service] = lambda: search_service

    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://testserver"
        ) as client:
            response = await client.post(
                "/api/v1/documents/search",
                headers={"X-Internal-User-ID": "user-a"},
                json={
                    "query": "FastAPI retrieval PostgreSQL pgvector",
                    "document_ids": [str(ingested.document.id)],
                    "limit": 5,
                },
            )
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    result = response.json()["results"]
    assert len(result) == 1
    assert result[0] == {
        "chunk_id": str(expected_match.chunk.id),
        "document_id": str(ingested.document.id),
        "filename": "profile.docx",
        "document_type": "docx",
        "chunk_index": 0,
        "text": "FastAPI retrieval\nPostgreSQL pgvector",
        "char_start": 0,
        "char_end": len("FastAPI retrieval\nPostgreSQL pgvector"),
        "score": pytest.approx(expected_match.score),
        "source_metadata": {"paragraph_count": 2},
    }
