from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Protocol

from docx import Document as DocxDocument
from pypdf import PdfReader


class UnsupportedDocumentTypeError(ValueError):
    """Raised when the uploaded file type has no registered parser."""


class DocumentExtractionError(ValueError):
    """Raised when a supported document cannot be read."""


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    metadata: dict[str, int | str] = field(default_factory=dict)


class DocumentParser(Protocol):
    def extract(self, content: bytes) -> ExtractedDocument: ...


class PlainTextParser:
    def extract(self, content: bytes) -> ExtractedDocument:
        try:
            return ExtractedDocument(text=content.decode("utf-8"))
        except UnicodeDecodeError as exc:
            raise DocumentExtractionError("Text documents must use UTF-8 encoding") from exc


class PdfParser:
    def extract(self, content: bytes) -> ExtractedDocument:
        try:
            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return ExtractedDocument(text=text, metadata={"page_count": len(reader.pages)})
        except Exception as exc:
            raise DocumentExtractionError("Unable to extract text from PDF") from exc


class DocxParser:
    def extract(self, content: bytes) -> ExtractedDocument:
        try:
            document = DocxDocument(BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            return ExtractedDocument(
                text=text, metadata={"paragraph_count": len(document.paragraphs)}
            )
        except Exception as exc:
            raise DocumentExtractionError("Unable to extract text from DOCX") from exc


class DocumentParserRegistry:
    """Resolve document parsers by trusted MIME type and filename extension."""

    _markdown_mime_types = {"text/markdown", "text/x-markdown"}
    _docx_mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def __init__(self) -> None:
        self._text_parser = PlainTextParser()
        self._pdf_parser = PdfParser()
        self._docx_parser = DocxParser()

    def get_parser(self, filename: str, mime_type: str) -> tuple[str, DocumentParser]:
        extension = Path(filename).suffix.lower()
        if mime_type == "application/pdf" or extension == ".pdf":
            return "pdf", self._pdf_parser
        if mime_type == self._docx_mime_type or extension == ".docx":
            return "docx", self._docx_parser
        if mime_type in self._markdown_mime_types or extension in {".md", ".markdown"}:
            return "markdown", self._text_parser
        if mime_type == "text/plain" or extension == ".txt":
            return "txt", self._text_parser
        raise UnsupportedDocumentTypeError(f"Unsupported document type: {mime_type or extension}")
